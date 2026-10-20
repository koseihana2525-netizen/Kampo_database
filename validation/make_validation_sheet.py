# -*- coding: utf-8 -*-
"""
make_validation_sheet.py - 第2評価者による検証用シートを作る

層（セット × 4象限）ごとに抄録を無作為抽出し、象限（判定器の結果）を伏せた評価シートを作る。
評価者は、題名と抄録だけを読んで「処方名の言及」と「漢方診断の記載」を判定する。
判定器の結果は別ファイル（answer_key.csv）に出す。評価が終わるまで評価者に見せない。

Usage
  python validation/make_validation_sheet.py --db PATH/integrated_db_japan.json --out DIR [--seed 2026]
出力（--out）
  validation_sheet.xlsx   評価者に渡す（象限は書かれていない）。評価者ごとに1部ずつ複製して使う
  answer_key.csv          id, セット, 層（象限）, 判定器が見つけた処方名・概念語、層の母集団の件数
  sampling_plan.json      抽出の条件（シード、層ごとの母集団と標本の件数）
標準ライブラリ + openpyxl。
"""

import argparse
import csv
import json
import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import gap_classifier as G
from analysis_set import analysis_records

# 層ごとの標本の件数。Q1 と Q3 は処方言及論文の分母なので多めに、Q2 と Q4 は処方名の取りこぼしの確認用
SAMPLE_SIZES = {("jp", "Q1_both"): 50, ("jp", "Q2_cognition_only"): 30, ("jp", "Q3_formula_only"): 60, ("jp", "Q4_neither"): 30,
                ("pm", "Q1_both"): 60, ("pm", "Q2_cognition_only"): 30, ("pm", "Q3_formula_only"): 60, ("pm", "Q4_neither"): 30}
SET_NAME = {"kampo": "jp", "pubmed_kampo": "pm"}
QUAD_SHORT = {"Q1_both": "Q1", "Q2_cognition_only": "Q2", "Q3_formula_only": "Q3", "Q4_neither": "Q4"}


def stratify(articles, variant="primary"):
    strata = {}
    for a in analysis_records(articles):
        c = G.classify({"title": a["title"], "abstract": a["abstract"], "lang": a["lang"]}, variant)
        key = (SET_NAME[a["source"]], c["quad_conservative"])
        strata.setdefault(key, []).append((a, c))
    return strata


def draw(strata, seed):
    """層ごとに無作為抽出し、各層を A・B の2つの束に交互に割り当てる。返り値: 行のリスト（束ごとにシャッフル済み）"""
    rng = random.Random(seed)
    rows = []
    for key in sorted(strata):
        pop = strata[key]
        n = min(SAMPLE_SIZES[key], len(pop))
        picked = rng.sample(pop, n)
        for i, (a, c) in enumerate(picked):
            rows.append(dict(batch="AB"[i % 2], stratum=key, article=a, cls=c, population=len(pop), sample=n))
    out = []
    for batch in "AB":
        part = [r for r in rows if r["batch"] == batch]
        rng.shuffle(part)
        out += part
    for i, r in enumerate(out, 1):
        r["id"] = f"V{i:03d}"
    return out


INSTRUCTIONS = [
    "【検証用シートの使い方】",
    "",
    "目的: 漢方関連論文の抄録に、処方名と漢方の診断（証・病態・診察所見）が書かれているかを、人の目で判定していただき、",
    "自動判定（辞書による判定）の正確さを調べます。自動判定の結果は、このシートには書かれていません。",
    "",
    "お願いすること:",
    "  1. 「評価シート」の各行について、題名と抄録だけを読み、C列とD列を選んでください（プルダウン）。",
    "  2. 迷った点や、判断の根拠になった語があれば、E列にメモしてください（任意）。",
    "  3. 他の評価者の回答や、自動判定の結果は見ないでください。独立した判定が必要です。",
    "  4. {N}件あります。AとBの2つの束（F列）に分かれています。Aだけ先に終えていただいても構いません。順番は自由です。",
    "  5. 目安は1件1〜2分です。途中で保存して、何回かに分けて構いません。",
    "",
    "判定する項目:",
    "  C列「処方名の言及」: 漢方の処方名（葛根湯、Shosaikoto、TJ-54 など）が書かれているか。",
    "      有 / 無 / 不明。生薬名だけ（黄連、Ephedra など）は『無』。複数の処方を併用していても『有』。",
    "  D列「漢方診断の記載」: 次のいずれかが、この論文の対象（患者・症例・処方の選び方）の説明として書かれているか。",
    "      ・証（虚証・実証、○○湯証、随証、方証相対 など）",
    "      ・漢方の病態（気血水の異常、寒熱、表裏、陰陽、六病位、臓腑の虚 など）",
    "      ・漢方の診察所見（腹診・腹証、舌診、脈診 など）",
    "      有: 上のどれかが、対象の説明・処方選択の根拠として書かれている。",
    "      背景のみ: 処方の伝統的な解説や一般論として語が出るだけで、この研究の対象の説明ではない。",
    "      無: 書かれていない。",
    "      不明: 判断できない（抄録が短すぎる、など）。",
    "  処方名が書かれていない抄録でも、D列は判定してください。",
    "",
    "迷ったときの考え方: 『この抄録を読んだ漢方医が、診断の情報を得られるか』で判断してください。",
    "処方名から証を推測できる場合（例: 当帰芍薬散）でも、抄録に診断が書かれていなければ『無』です。",
    "",
    "お返事: 評価が終わったら、このファイルをそのまま返送してください（ファイル名に評価者の名前を付けてください）。",
]

CRITERIA = [
    ["判定", "意味", "例"],
    ["処方名: 有", "漢方処方の名前が書かれている", "「抑肝散を投与した」「Shosaikoto」「TJ-54」「半夏瀉心湯と六君子湯を併用」"],
    ["処方名: 無", "処方名がない（生薬名、成分名、一般的な『漢方薬』のみ）", "「黄連の抗炎症作用」「Kampo medicine was effective」"],
    ["診断: 有", "証・漢方の病態・診察所見が、対象の説明や処方選択の根拠になっている", "「虚証の患者に補中益気湯を選択」「腹診で胸脇苦満を認め」「瘀血の所見」「kyosho patients」"],
    ["診断: 背景のみ", "処方の伝統的解説・一般論で語が出るだけ", "「本方は古来、血の道症に用いられてきた」「traditionally used to improve blood stasis (oketsu)」"],
    ["診断: 無", "上のどれも書かれていない", "「うつ病に対する抑肝散の効果を検討した」"],
    ["不明", "判断できない", "抄録が数行しかない、など"],
]


def build_workbook(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "使い方"
    for i, line in enumerate(INSTRUCTIONS, 1):
        ws0.cell(row=i, column=1, value=line.replace("{N}", str(len(rows))))
    ws0.column_dimensions["A"].width = 130
    ws1 = wb.create_sheet("判定基準")
    for r, line in enumerate(CRITERIA, 1):
        for c, v in enumerate(line, 1):
            cell = ws1.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r == 1:
                cell.font = Font(bold=True)
    for col, w in zip("ABC", (18, 60, 80)):
        ws1.column_dimensions[col].width = w
    ws = wb.create_sheet("評価シート")
    head = ["ID", "言語", "C 処方名の言及", "D 漢方診断の記載", "E メモ（任意）", "束", "年", "題名", "抄録"]
    fill = PatternFill("solid", fgColor="E8F0FE")
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = fill
    for i, r in enumerate(rows, 2):
        a = r["article"]
        vals = [r["id"], "日本語" if a["lang"] == "ja" else "English", None, None, None, r["batch"], a.get("year"), a["title"], a["abstract"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    dv1 = DataValidation(type="list", formula1='"有,無,不明"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"有,背景のみ,無,不明"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"C2:C{len(rows) + 1}")
    dv2.add(f"D2:D{len(rows) + 1}")
    for col, w in zip("ABCDEFGHI", (7, 9, 14, 16, 30, 5, 7, 45, 100)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="第2評価者による検証用シートを作る")
    ap.add_argument("--db", required=True, help="integrated_db_japan.json のパス")
    ap.add_argument("--out", required=True, help="出力先のディレクトリ")
    ap.add_argument("--seed", type=int, default=2026)
    args = ap.parse_args()
    raw = json.load(open(args.db, encoding="utf-8"))
    articles = raw["articles"] if isinstance(raw, dict) else raw
    strata = stratify(articles)
    rows = draw(strata, args.seed)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    build_workbook(rows, out / "validation_sheet.xlsx")
    with open(out / "answer_key.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "batch", "set", "stratum", "stratum_population", "stratum_sample", "classifier_formula", "classifier_formulas", "classifier_concept", "classifier_terms", "article_id"])
        for r in rows:
            c = r["cls"]
            w.writerow([r["id"], r["batch"], r["stratum"][0], QUAD_SHORT[r["stratum"][1]], r["population"], r["sample"], int(c["formula"]), "|".join(c["formulas"][:5]),
                        int(bool(c["conservative"])), "|".join(sorted({t for t, _ in c["conservative"]})), r["article"]["id"]])
    plan = {"seed": args.seed, "variant": "primary", "strata": {f"{k[0]}:{QUAD_SHORT[k[1]]}": {"population": len(v), "sample": min(SAMPLE_SIZES[k], len(v))} for k, v in sorted(strata.items())}}
    json.dump(plan, open(out / "sampling_plan.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"{len(rows)} 件を抽出: {out}")
    for k, v in plan["strata"].items():
        print(f"  {k}: 母集団 {v['population']}、標本 {v['sample']}")


if __name__ == "__main__":
    main()

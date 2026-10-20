# -*- coding: utf-8 -*-
"""検証用シート（validation/）と解析対象の選び方（analysis_set.py）のテスト。合成の短い文だけを使う。"""

import csv
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "validation"))

import analysis_set as A

try:
    import openpyxl  # noqa: F401
    HAVE_XLSX = True
except ImportError:
    HAVE_XLSX = False


def art(i, source, lang, title, abstract, pub_types=None):
    return {"id": f"{source}_{i}", "source": source, "lang": lang, "year": "2010", "title": title, "abstract": abstract, "pub_types": pub_types or []}


def synthetic():
    arts = []
    for i in range(12):
        arts.append(art(f"a{i}", "kampo", "ja", f"症例{i}", f"抑肝散を投与した。虚証の患者であった。{i}"))        # Q1
        arts.append(art(f"b{i}", "kampo", "ja", f"報告{i}", f"抑肝散を投与した。経過は良好であった。{i}"))      # Q3
        arts.append(art(f"c{i}", "kampo", "ja", f"概説{i}", f"気虚と血虚の病態を述べる。{i}"))               # Q2
        arts.append(art(f"d{i}", "kampo", "ja", f"検査{i}", f"検査値の推移を検討した。{i}"))                  # Q4
        arts.append(art(f"e{i}", "pubmed_kampo", "en", f"Case {i}", f"Yokukansan was selected according to Sho. Case {i}."))   # Q1
        arts.append(art(f"f{i}", "pubmed_kampo", "en", f"Trial {i}", f"Yokukansan was given to patients. Study {i}."))        # Q3
        arts.append(art(f"g{i}", "pubmed_kampo", "en", f"Theory {i}", f"The concept of qi deficiency is described. Note {i}."))  # Q2
        arts.append(art(f"h{i}", "pubmed_kampo", "en", f"Lab {i}", f"Serum levels were measured. Sample {i}."))               # Q4
    arts.append(art("err", "pubmed_kampo", "en", "Erratum", "Yokukansan was given.", ["Published Erratum"]))
    arts.append(art("noabs", "kampo", "ja", "抄録なし", ""))
    arts.append(art("other", "acupuncture", "ja", "鍼灸", "鍼治療を行った。"))
    return arts


class AnalysisSet(unittest.TestCase):
    def test_exclusions(self):
        arts = synthetic() + [art("dup", "pubmed_kampo", "en", "Dup", "Serum levels were measured. Sample 0.")]
        ex = A.exclusion_reasons(arts)
        self.assertEqual(ex["pubmed_kampo_err"], "published erratum")
        self.assertEqual(ex["pubmed_kampo_dup"], "duplicate abstract")
        self.assertNotIn("kampo_noabs", ex)           # 抄録なしは除外理由ではない（別に除かれる）
        ids = {a["id"] for a in A.analysis_records(arts)}
        self.assertNotIn("kampo_noabs", ids)
        self.assertNotIn("acupuncture_other", ids)
        self.assertNotIn("pubmed_kampo_err", ids)
        self.assertIn("pubmed_kampo_h0", ids)


@unittest.skipUnless(HAVE_XLSX, "openpyxl が無い")
class ValidationSheet(unittest.TestCase):
    def setUp(self):
        import make_validation_sheet as M
        import analyze_validation as V
        self.M, self.V = M, V
        self.sizes = dict(M.SAMPLE_SIZES)
        M.SAMPLE_SIZES.update({k: 100 for k in M.SAMPLE_SIZES})   # 母集団より大きい → 全数

    def tearDown(self):
        self.M.SAMPLE_SIZES.update(self.sizes)

    def test_sheet_hides_the_classification_and_key_has_it(self):
        from openpyxl import load_workbook
        rows = self.M.draw(self.M.stratify(synthetic()), seed=1)
        self.assertEqual(len(rows), 8 * 12)
        self.assertEqual(sorted(r["id"] for r in rows)[0], "V001")
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "s.xlsx"
            self.M.build_workbook(rows, p)
            wb = load_workbook(p)
            self.assertEqual(wb.sheetnames, ["使い方", "判定基準", "評価シート"])
            ws = wb["評価シート"]
            heads = [c.value for c in ws[1]]
            self.assertNotIn("stratum", " ".join(map(str, heads)))
            dump = " ".join(str(c.value) for row in ws.iter_rows() for c in row)
            self.assertNotIn("Q1_both", dump)
            self.assertEqual(ws.max_row, 1 + len(rows))
        self.assertEqual({r["batch"] for r in rows}, {"A", "B"})

    def test_oracle_raters_reproduce_the_observed_gap_and_kappa_is_one(self):
        from openpyxl import load_workbook
        M, V = self.M, self.V
        strata = M.stratify(synthetic())
        rows = M.draw(strata, seed=1)
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            M.build_workbook(rows, d / "s.xlsx")
            with open(d / "key.csv", "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(["id", "batch", "set", "stratum", "stratum_population", "stratum_sample", "classifier_formula", "classifier_formulas", "classifier_concept", "classifier_terms", "article_id"])
                for r in rows:
                    c = r["cls"]
                    w.writerow([r["id"], r["batch"], r["stratum"][0], M.QUAD_SHORT[r["stratum"][1]], r["population"], r["sample"], int(c["formula"]), "", int(bool(c["conservative"])), "", r["article"]["id"]])
            key = V.load_key(d / "key.csv")
            for name in ("x", "y"):
                wb = load_workbook(d / "s.xlsx")
                ws = wb["評価シート"]
                for row in ws.iter_rows(min_row=2):
                    k = key[row[0].value]
                    row[2].value = "有" if k["formula"] else "無"
                    row[3].value = "有" if k["concept"] else "無"
                wb.save(d / f"{name}.xlsx")
            raters = {n: V.load_rater(d / f"{n}.xlsx") for n in ("x", "y")}
            pairs = [(V.to_formula(raters["x"][k][0]), V.to_formula(raters["y"][k][0])) for k in key]
            self.assertAlmostEqual(V.kappa(pairs)[0], 1.0)
            items = [((r["set"], r["stratum"]), r["population"], dict(cf=r["formula"], cc=r["concept"], f=V.to_formula(raters["x"][k][0]), d=V.to_diag(raters["x"][k][1])))
                     for k, r in key.items() if r["set"] == "jp"]
            e = V.estimates(items)
            self.assertAlmostEqual(e["gap"], V.observed_gap(key, "jp"))
            self.assertAlmostEqual(e["f_ppv"], 1.0)
            text = "\n".join(V.report(key, raters))
            self.assertIn("補正後のギャップ", text)

    def test_kappa_and_wilson(self):
        V = self.V
        self.assertAlmostEqual(V.kappa([(1, 1), (0, 0), (1, 0), (0, 1)])[0], 0.0)
        lo, hi = V.wilson(5, 10)
        self.assertLess(lo, 0.5)
        self.assertGreater(hi, 0.5)


if __name__ == "__main__":
    unittest.main()
